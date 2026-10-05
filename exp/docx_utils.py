"""Small helpers for building the article .docx in the original's format
(Times New Roman 12 pt, justified text, centred captions)."""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Times New Roman"


def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.0
    return doc


def para(doc, text="", bold=False, italic=False, align="justify", indent=True, size=None, space_after=0):
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    if indent and align == "justify":
        p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.space_after = Pt(space_after)
    _runs(p, text, bold, italic, size)
    return p


def _runs(p, text, bold=False, italic=False, size=None):
    """Plain text with **bold** fragments."""
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run(part)
        r.bold = bold or (i % 2 == 1)
        r.italic = italic
        r.font.name = FONT
        if size:
            r.font.size = Pt(size)


def heading(doc, text):
    p = para(doc, text, bold=True, align="center", indent=False)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _runs(p, text)
    return p


def numbered(doc, text, n, justify=True):
    """Explicitly numbered item (Word's List Number style would continue numbering across lists)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.75)
    _runs(p, f"{n}.	{text}")
    return p


def figure(doc, path, caption, width_cm=17.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(width_cm))
    para(doc, caption, align="center", indent=False, size=11, space_after=6)


def table(doc, caption, header, rows, size=10, col_widths=None):
    para(doc, caption, align="right", indent=False, italic=True, size=11).paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(t.rows[0].cells, header):
        cell.text = ""
        _runs(cell.paragraphs[0], text, bold=True, size=size)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = ""
            _runs(cell.paragraphs[0], str(text), size=size)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    # keep small tables on one page: every row except the last stays with the next one
    for row in t.rows[:-1]:
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.keep_with_next = True
    if col_widths:
        for row in t.rows:
            for cell, w in zip(row.cells, col_widths):
                cell.width = Cm(w)
    para(doc, "", space_after=4)
    return t
